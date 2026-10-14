"""Check that the generated dataset is internally consistent.

Usage: uv run python src/validate.py   (run src/build.py first)
Reads the generated files back from disk instead of trusting the generator.
"""
from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

import comms as C
import scenario as S

ROOT = Path(__file__).resolve().parent.parent
DS = ROOT / "dataset"
errors: list[str] = []


def check(cond, msg):
    if not cond:
        errors.append(msg)


def read_csv(path):
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def line_totals(rows):
    tot = defaultdict(Decimal)
    for r in rows:
        line = S.ACCOUNTS[r["account"]][1]
        tot[line] += Decimal(r["credit"]) - Decimal(r["debit"])
    return {k: v * S.LINE_SIGN[k] for k, v in tot.items()}


def ebitda(values):
    return sum(Decimal(values[k]) * S.LINE_SIGN[k] for k in values)


# 1. General ledger exports reconcile to the actuals, line by line.
gl = {}
for m in S.MONTHS:
    rows = read_csv(DS / "erp-exports" / f"gl_{m}.csv")
    gl[m] = rows
    ids = [r["entry_id"] for r in rows]
    check(len(ids) == len(set(ids)), f"{m}: duplicate entry_id in valid export")
    tot = line_totals(rows)
    for code, _, _ in S.LINES:
        want = Decimal(S.ACTUALS[m][code])
        check(tot.get(code, Decimal(0)) == want, f"{m} {code}: GL {tot.get(code)} != actual {want}")
    check(all(r["period"] == m and r["posting_date"].startswith(m) for r in rows), f"{m}: period/date mismatch")
    check(m == "2026-03" or not any(r["account"] == "4020" for r in rows), f"{m}: account 4020 before March")
check(not any(r["counterparty"] == "Belvedere Hotel Group" and r["account"] == "4000" for r in gl["2026-03"]),
      "Belvedere revenue must be absent in March (order slipped)")

# 2. The invalid March export fails in exactly the two planned ways.
inv = read_csv(DS / "erp-exports" / "gl_2026-03_invalid.csv")
dups = [k for k, n in Counter(r["entry_id"] for r in inv).items() if n > 1]
bad = []
for r in inv:
    for col in ("debit", "credit"):
        try:
            Decimal(r[col])
            if not re.fullmatch(r"\d+\.\d{2}", r[col]):
                bad.append(r["entry_id"])
        except InvalidOperation:
            bad.append(r["entry_id"])
check(len(dups) == 1, f"invalid export: expected 1 duplicate entry_id, got {dups}")
check(len(bad) == 1, f"invalid export: expected 1 malformed amount, got {bad}")
check(dups and next(r for r in inv if r["entry_id"] == dups[0])["posting_date"] == "2026-03-14",
      "duplicate should sit on the split-export boundary day (14 March)")

# 3. Answer key: drivers add up, every material line is explained, evidence exists.
key = json.loads((DS / "answer-key" / "march_2026_variances.json").read_text())
ids = {e["id"] for e in C.EMAILS} | {c["id"] for c in C.CHATS} | {m["id"] for m in C.MEETINGS} | {d["id"] for d in C.DOCUMENTS}
docs_march = {r["document_ref"] for r in gl["2026-03"]}
for line in key["lines"]:
    if line["above_threshold"]:
        check(line["drivers"], f"answer key: {line['line_code']} above threshold without drivers")
    check(not line["above_threshold"] or line.get("repeats_next_month"), f"answer key: {line['line_code']} missing repeats_next_month")
    if line["drivers"]:
        s = sum(d["amount"] for d in line["drivers"])
        check(s == line["variance_favourable_positive"], f"answer key: {line['line_code']} drivers {s} != variance {line['variance_favourable_positive']}")
    for d in line["drivers"]:
        for ev in d["evidence"]:
            if ev.startswith("GL:"):
                check(ev[3:] in docs_march, f"evidence {ev} not in March GL")
            elif ev.startswith("COMMENTARY:"):
                check((DS / "commentary" / f"commentary_{ev.split(':')[1]}.md").exists(), f"evidence {ev} missing")
            else:
                check(ev in ids, f"evidence {ev} unknown")
check(key["ebitda_actual"] == ebitda(S.ACTUALS["2026-03"]), "answer key EBITDA mismatch")
# Recompute every line independently: actual from the GL file, budget from the budget CSV.
budget_csv = {(r["month"], r["line_code"]): Decimal(r["amount"]) for r in read_csv(DS / "budget" / "budget_2026.csv")}
pnl_csv = {(r["month"], r["line_code"]): r for r in read_csv(DS / "answer-key" / "pnl_by_month.csv")}


def expected_line(month, code, gl_tot):
    act = gl_tot.get(code, Decimal(0))
    bud = budget_csv[(month, code)]
    var = (act - bud) * S.LINE_SIGN[code]
    pct = None if bud == 0 else round(float(var / bud * 100), 1)
    return act, bud, var, pct, abs(var) >= S.COMPANY["variance_threshold"]


for m in S.MONTHS:
    tot = line_totals(gl[m])
    for code, _, _ in S.LINES:
        act, bud, var, pct, above = expected_line(m, code, tot)
        r = pnl_csv[(m, code)]
        got = (Decimal(r["actual"]), Decimal(r["budget"]), Decimal(r["variance_favourable_positive"]),
               None if r["variance_pct"] == "" else float(r["variance_pct"]), r["above_threshold"] == "True")
        check(got == (act, bud, var, pct, above), f"pnl_by_month {m} {code}: {got} != {(act, bud, var, pct, above)}")
tot = line_totals(gl["2026-03"])
for line in key["lines"]:
    act, bud, var, pct, above = expected_line("2026-03", line["line_code"], tot)
    got = (Decimal(line["actual"]), Decimal(line["budget"]), Decimal(line["variance_favourable_positive"]),
           line["variance_pct"], line["above_threshold"])
    check(got == (act, bud, var, pct, above), f"answer key {line['line_code']}: {got} != {(act, bud, var, pct, above)}")
material = sorted(l["line_code"] for l in key["lines"] if l["above_threshold"])
check(material == sorted(S.MARCH_DRIVERS), f"material lines {material} != drivers {sorted(S.MARCH_DRIVERS)}")

# 4. Story texts agree with the figures.
feb_e = ebitda(S.ACTUALS["2026-02"])
e05 = next(e for e in C.EMAILS if e["id"] == "E05")["body"]
check(f"{feb_e / 1000:.1f}k" in e05, f"E05 should state February EBITDA {feb_e}")
feb_timber = (S.ACTUALS["2026-02"]["MAT_TIMBER"] - S.BUDGET["2026-02"]["MAT_TIMBER"]) / 1000
check(f"-{feb_timber:.1f}k" in e05, "E05 timber variance mismatch")
# The freight surcharges must stay unexplained: no human message gives a reason.
human = [e["body"] for e in C.EMAILS] + [c["text"] for c in C.CHATS] + [t for m in C.MEETINGS for _, t in m["lines"]]
check(not any("express" in t.lower() and "because" in t.lower() for t in human), "freight surcharge must stay unexplained")
check(not any(re.search(r"\bexpress\b", t, re.I) for t in [e["body"] for e in C.EMAILS] + [c["text"] for c in C.CHATS]),
      "no email or chat may mention express shipments")
# The temporary labour reason is only said in the production meeting.
written = [e["body"] for e in C.EMAILS] + [c["text"] for c in C.CHATS] + [d["note"] for d in C.DOCUMENTS]
check(not any("Harbour" in t and re.search(r"temporar|operative|StaffBridge", t, re.I) for t in written),
      "the link temps <-> Harbour City must only be spoken in meeting M01")
check(any("Harbour" in t and "operatives" in t for _, t in next(m for m in C.MEETINGS if m["id"] == "M01")["lines"]),
      "meeting M01 must link the operatives to the Harbour City order")

# 5. screenpipe-style records: starter-compatible shape.
for p in sorted((DS / "screenpipe-records").glob("*.json")):
    data = json.loads(p.read_text())
    check(isinstance(data, list) and data, f"{p.name}: not a non-empty array")
    check(p.stat().st_size < 2_000_000, f"{p.name}: over 2 MB")
    rids = [r.get("id") for r in data]
    check(len(rids) == len(set(rids)), f"{p.name}: duplicate ids")
    for r in data:
        check(set(r) == {"id", "timestamp", "app", "text"}, f"{p.name}: bad keys {set(r)}")
        check(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", r["timestamp"]) is not None, f"{p.name}: bad timestamp {r['timestamp']}")
runs = json.loads((DS / "screenpipe-records" / "task-runs.json").read_text())
for r in runs:
    check(re.fullmatch(r"Task: .+\nRun: .+\nMinutes: \d+", r["text"]) is not None, f"task-runs: bad annotation {r['id']}")
for name in ("handoff.json", "escalation.json"):
    for r in json.loads((DS / "screenpipe-records" / name).read_text()):
        for ln in r["text"].splitlines():
            if re.match(r"^(Decision|Blocker|Action|Symptom|Attempt|Result|Question)\b", ln):
                check(re.match(r"^(Decision|Blocker|Symptom|Attempt|Result|Question): \S|^Action \([^)]+\): \S", ln) is not None, f"{name}: bad line '{ln}'")

# 6. Workbooks and decks open, with the expected structure.
for m in S.MONTHS:
    wb = load_workbook(DS / "workbooks" / f"closing_workbook_{m}.xlsx")
    check(wb.sheetnames == ["README", "GL_import", "Mapping", "PnL_vs_Budget", "Commentary"], f"workbook {m}: sheets {wb.sheetnames}")
    check(wb["GL_import"].max_row == len(gl[m]) + 1, f"workbook {m}: GL row count")
    accts = {str(r[0].value) for r in wb["Mapping"].iter_rows(min_row=2)}
    check("4020" not in accts, f"workbook {m}: mapping must predate account 4020")
    check((wb["Commentary"].max_row > 1) == (m != "2026-03"), f"workbook {m}: commentary presence")
check((ROOT / "docs" / "pitch-deck.pptx").exists() and len(Presentation(ROOT / "docs" / "pitch-deck.pptx").slides) == 12,
      "docs/pitch-deck.pptx missing or not 12 slides (run src/pitch_deck.py)")
for f in (DS / "presentations").glob("*.pptx"):
    check(len(Presentation(f).slides) >= 3, f"{f.name}: too few slides")

# 7. Viewer data covers everything.
js = (DS / "viewer" / "data.js").read_text()
check(js.startswith("window.DATASET = "), "viewer data.js prefix")
vd = json.loads(js[len("window.DATASET = "):].rstrip().rstrip(";"))
check(len(vd["emails"]) == len(C.EMAILS) and len(vd["chats"]) == len(C.CHATS) and len(vd["meetings"]) == len(C.MEETINGS), "viewer counts")
check(set(vd["gl"]) == set(S.MONTHS), "viewer GL months")

# 8. Fictional email domains only.
all_text = "\n".join(p.read_text(errors="ignore") for p in DS.rglob("*") if p.suffix in {".eml", ".md", ".json", ".csv", ".html", ".js"})
emails_found = set(re.findall(r"[\w.+-]+@([\w.-]+)", all_text))
check(all(d.endswith(".example") for d in emails_found), f"non-.example email domains: {sorted(d for d in emails_found if not d.endswith('.example'))}")

if errors:
    print(f"{len(errors)} problem(s):")
    for e in errors:
        print(" -", e)
    sys.exit(1)
print(f"OK - dataset consistent ({sum(len(v) for v in gl.values())} GL entries, {len(C.EMAILS)} emails, "
      f"{len(C.CHATS)} chat messages, {len(C.MEETINGS)} meetings, {len(C.DOCUMENTS)} documents)")
