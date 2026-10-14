"""Build the hackathon pitch deck as an editable PowerPoint file.

Usage: uv run python src/pitch_deck.py   ->  docs/pitch-deck.pptx
Fonts: Arial only, so the deck looks the same on any laptop.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from build import FIXED_TIME, normalize_zip

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "pitch-deck.pptx"

FONT = "Arial"
DARK, LIGHT, SAND = "14202E", "F6F4EF", "ECE8DF"
ORANGE, ORANGE_LIGHT, BLUE = "B85A24", "F0A36F", "2F6DA8"
BODY, MUTED, ON_DARK, ON_DARK_MUTED, RULE = "3D4A5A", "5F6B78", "C9D3DD", "9FB0C0", "D8D2C4"

# Layout grid in pixels of a 1920x1080 canvas (1 px = 6350 EMU on a 13.33 in slide).
PX = 6350
M = 128  # margin
W = 1920 - 2 * M


def rgb(h):
    return RGBColor.from_string(h)


def new_slide(prs, bg, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(bg)
    s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: a string, or a list of paragraphs; a paragraph is a string or a list of (text, bold) pairs."""
    box = s.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, b in ([(para, bold)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size / 2), b
            r.font.color.rgb = rgb(color)
    return box


def card(s, x, y, w, h, fill, radius=True):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = min(0.5, 16 / min(w, h))
    return shape


def footer(s, txt, color):
    text(s, M, 1080 - 64 - 30, W, 30, txt, 24, color)


def title(s, txt, color=DARK, y=M, size=72, w=W):
    return text(s, M, y, w, int(size * 1.2 * 2), txt, size, color, bold=True)


def table(s, x, y, w, rows, widths, size=24, header_fill=SAND, row_h=None):
    row_h = row_h or int(size * 2.2)
    shape = s.shapes.add_table(len(rows), len(rows[0]), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(row_h * len(rows) * PX))
    t = shape.table
    for j, frac in enumerate(widths):
        t.columns[j].width = Emu(int(w * frac) * PX)
    for i, row in enumerate(rows):
        t.rows[i].height = Emu(row_h * PX)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else (LIGHT if i % 2 else "EFEBE3"))
            cell.margin_left = cell.margin_right = Emu(12 * PX)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            bold = i == 0 or (isinstance(val, tuple) and val[1])
            txt = val[0] if isinstance(val, tuple) else val
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size / 2), bold
            r.font.color.rgb = rgb(DARK)
            if isinstance(txt, str) and txt[:1] in "+-" and txt[1:2].isdigit():
                p.alignment = PP_ALIGN.RIGHT
    return shape


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)

    # 1. Cover
    s = new_slide(prs, DARK, "Open with the problem, not the tech. One sentence: every month a financial controller explains why the numbers moved, and the explanation has already crossed their screen.")
    text(s, M, 300, W, 50, "42 PARIS X SCREENPIPE - BUILD AGENTS THAT REMEMBER", 32, ORANGE_LIGHT, bold=True)
    text(s, M, 380, W, 150, "Close Memory", 120, LIGHT, bold=True)
    text(s, M, 560, 1400, 130, "An agent that closes the month with screenpipe as its memory, and explains every variance with its source.", 44, ON_DARK)
    footer(s, "Track 2 - Teach an agent your workflow - Fictional company: Pipeline Furniture Ltd", ON_DARK_MUTED)

    # 2. Problem
    s = new_slide(prs, LIGHT, "Three sources, three channels. The information exists; it is scattered across a month of screen time.")
    text(s, M, M, W, 45, "THE PROBLEM", 32, ORANGE, bold=True)
    title(s, "Every month, the controller explains why the numbers moved. The answer already crossed their screen.", y=190, size=60)
    cw = (W - 2 * 32) // 3
    for i, (h, q) in enumerate([("Email", '"Prices up 8% from 1 February."'),
                                ("Chat", '"Belvedere pushed the 45k order to April."'),
                                ("Meeting", '"Three temps for the rush order, about 14k."')]):
        x = M + i * (cw + 32)
        card(s, x, 470, cw, 220, SAND)
        text(s, x + 40, 500, cw - 80, 55, h, 44, ORANGE, bold=True)
        text(s, x + 40, 575, cw - 80, 100, q, 32, BODY)
    text(s, M, 740, 1550, 100, "At month end it all has to be found again, by hand, under deadline. Searching takes the time that analysis should get.", 32, BODY)
    footer(s, "Quotes from the fictional Pipeline Furniture dataset", MUTED)

    # 3. Pain points
    s = new_slide(prs, LIGHT, "Do not read the table. Point at the six Demo rows: that is what the jury will see. The rest is the roadmap; the dataset already supports two of them.")
    title(s, "Eleven pain points, and where we answer", size=64)
    rows = [["Pain point", "Track", "How the agent answers", "Status"],
            ["Explanation crossed the screen, dug up again", "1", "Finds the email, chat or meeting", ("Demo", True)],
            ["Same variances re-researched every month", "1", "Recalls last month's explanation", ("Demo", True)],
            ["Some facts are only said in meetings", "1", "Searches the audio transcript", ("Demo", True)],
            ["Exports and spreadsheet work repeated monthly", "1, 2", "Replays approved steps with code", ("Demo", True)],
            ["Close procedure only in someone's head", "2", "Drafts it from one observed close", ("Demo", True)],
            ["Commentary starts from a blank page", "2", "Drafts it, every sentence sourced", ("Demo", True)],
            ["Controls skipped under deadline", "2", "Flags steps missing vs procedure", "Data ready"],
            ["Chasing answers colleagues already gave", "2", "Checks history, drafts the question", "Data ready"],
            ['"Where does this number come from?"', "1", "Rebuilds the evidence trail", "Vision"],
            ["Leave or departure, replacement is lost", "3", "Sourced handoff and guide", "Vision"],
            ["Procedure goes out of date", "3", "Proposes changes, re-approval", "Vision"]]
    table(s, M, 230, W, rows, [0.44, 0.08, 0.34, 0.14], size=24, row_h=50)
    footer(s, "Demo = shown live. Data ready = supported by our dataset. Vision = after the hackathon.", MUTED)

    # 4. Tracks
    s = new_slide(prs, SAND, "One project, one track, chosen Friday 19:15. Track 2 is the demo itself. Tracks 1 and 3 become the roadmap slide.")
    title(s, "Three tracks. We enter track 2.")
    cards = [("TRACK 1", "Find what to automate", "The slowest step of the close is searching for explanations.", "screenpipe already maps workflows: we would compete with it.", False),
             ("TRACK 2 - OUR TRACK", "Teach an agent your workflow", "The agent watches the February close, learns it, then closes March on its own, from data to sourced commentary.", "Exactly the brief: recorded work, then a real task done end to end.", True),
             ("TRACK 3", "Make it work for a team", "Handoff to a colleague on leave, onboarding, approved changes.", "Needs teammate sharing we would build from scratch: our vision.", False)]
    widths = [500, W - 2 * 500 - 2 * 32, 500]
    x = M
    for (eyebrow, h, body, note, ours), cw in zip(cards, widths):
        card(s, x, 290, cw, 620, DARK if ours else LIGHT)
        text(s, x + 40, 330, cw - 80, 35, eyebrow, 24, ORANGE_LIGHT if ours else MUTED, bold=True)
        text(s, x + 40, 380, cw - 80, 120, h, 44, LIGHT if ours else DARK, bold=True)
        text(s, x + 40, 520, cw - 80, 200, body, 28, ON_DARK if ours else BODY)
        text(s, x + 40, 760, cw - 80, 120, note, 24, ON_DARK_MUTED if ours else MUTED)
        x += cw + 32

    # 5. How it works
    s = new_slide(prs, LIGHT, "Steps 1 and 2 happen once, in February. Steps 3 to 5 are the March close the agent runs on its own. Stress step 3: the numbers never come from the model.")
    title(s, "How it works")
    steps = [("1", "Capture", "screenpipe records screen and meeting audio, locally"),
             ("2", "Learn", "One observed close becomes a procedure, each step with evidence"),
             ("3", "Compute", "Code computes the variances and rejects bad exports"),
             ("4", "Explain", "The agent links each variance to its email, chat or meeting"),
             ("5", "Approve", "The controller checks the sourced commentary and signs off")]
    gap = 56
    cw = (W - 4 * gap) // 5
    for i, (n, h, body) in enumerate(steps):
        x = M + i * (cw + gap)
        last = i == 4
        card(s, x, 330, cw, 360, DARK if last else SAND)
        text(s, x + 28, 360, cw - 56, 35, n, 24, ORANGE_LIGHT if last else ORANGE, bold=True)
        text(s, x + 28, 405, cw - 56, 50, h, 36, LIGHT if last else DARK, bold=True)
        text(s, x + 28, 470, cw - 56, 200, body, 24, ON_DARK if last else BODY)
        if not last:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu((x + cw + 10) * PX), Emu(490 * PX), Emu((gap - 20) * PX), Emu(36 * PX))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = rgb(BLUE)
            arrow.line.fill.background()
    text(s, M, 760, 1550, 100, "Built on screenpipe's local API and MCP, and on the official hackathon starter: same review desk, same evidence rules.", 32, BODY)

    # 6. Demo
    s = new_slide(prs, LIGHT, "Every line is a different kind of test: two stories netted in one number, a new account, memory, audio, no evidence at all, a one-off, a timing effect. The answer key is in this repository.")
    title(s, "The demo: closing March", size=64)
    text(s, M, 220, W, 45, [[("EBITDA 276.9k vs budget 324.0k: ", False), ("-47.1k", True), (". Seven lines need an explanation.", False)]], 32, BODY)
    rows = [["Line", "Variance", "Cause the agent must find", "Where it was"],
            ["Revenue - Contract", "-17,000", "Belvedere -45k slipped to April, Harbour City +28k rush", "Email, chat, meeting"],
            ["Revenue - Online", "+6,400", "New web shop, account not yet mapped", "Email, chat"],
            ["Timber and panels", "-12,400", "Supplier +8%, same cause as February", "Last month's commentary"],
            ["Temporary labour", "-14,000", "Three temps for the rush order", "Meeting audio only"],
            ["Freight out", "-6,200", "Express surcharges: no reason anywhere", "Nowhere: ask Sofia"],
            ["Energy", "+5,600", "One-off credit note", "Email, document"],
            ["Professional fees", "-8,000", "Invoiced in March, budgeted in April", "Emails"]]
    table(s, M, 300, W, rows, [0.21, 0.11, 0.43, 0.25], size=24, row_h=66)
    footer(s, "Fictional company, 163 ledger entries, 19 emails, 12 chat messages, 3 meetings. Positive = favourable.", MUTED)

    # 7. Five moments
    s = new_slide(prs, DARK, "These five map to the demo script. If time runs short live, keep 2, 3 and 4: audio, memory, honesty. They are the hackathon's theme.")
    title(s, "Five moments you will see", color=LIGHT)
    moments = [("One number, two stories.", " -17k on contracts is really -45k late and +28k extra."),
               ("It heard the meeting.", " The temps were only ever mentioned out loud."),
               ("It remembers.", ' "Same cause as February. Still true? Also +3% more from May."'),
               ("It knows what it does not know.", " No evidence for freight: it drafts a question instead of a guess."),
               ("It refuses bad data.", " A duplicated entry and a malformed amount stop the close before any figure.")]
    for i, (b, rest) in enumerate(moments):
        y = 270 + i * 120
        text(s, M, y, 70, 60, str(i + 1), 44, ORANGE_LIGHT, bold=True)
        text(s, M + 90, y + 4, W - 90, 100, [[(b, True), (rest, False)]], 32, "E3E8EE")

    # 8. Trust
    s = new_slide(prs, SAND, "Answers the jury's first question before they ask it: 'you record a finance person's whole screen?' The privacy controls are screenpipe's own: app and window filters, password exclusion, per-pipe permissions enforced by the system.")
    title(s, "Code computes. The model explains. A human approves.", w=1500)
    items = [("Numbers", "Every figure comes from tested code. The model never calculates."),
             ("Evidence", "Every explanation shows its exact quote and source. No source, no claim."),
             ("Approval", "The procedure and the commentary are drafts until the controller signs off."),
             ("Privacy", "Local by default, chosen apps only, passwords excluded, per-agent data permissions.")]
    cw = (W - 32) // 2
    for i, (h, body) in enumerate(items):
        x = M + (i % 2) * (cw + 32)
        y = 380 + (i // 2) * 270
        card(s, x, y, cw, 240, LIGHT)
        text(s, x + 40, y + 36, cw - 80, 50, h, 36, BLUE, bold=True)
        text(s, x + 40, y + 100, cw - 80, 120, body, 28, BODY)

    # 9. Why screenpipe could ship it
    s = new_slide(prs, LIGHT, "The prize is a paid contract to integrate the hack into the product. This slide is the business case. Do not claim hours saved: we have not measured them.")
    title(s, "Why screenpipe could ship it")
    reasons = [("It is your thesis", "Find work worth automating and give agents context, applied to a job with a clear owner and deliverable."),
               ("It extends your starter", "Review desk, exact quotes, frozen approvals, replay with fixed code: we follow the same rules."),
               ("It recurs everywhere", "Every company closes its books every month. Same pain, same shape, every time."),
               ("It is measurable", 'An answer key scores each variance: right cause, right source, right "repeats next month".')]
    for i, (h, body) in enumerate(reasons):
        y = 290 + i * 150
        line = s.shapes.add_connector(1, Emu(M * PX), Emu(y * PX), Emu((M + W) * PX), Emu(y * PX))
        line.line.color.rgb = rgb(RULE)
        line.line.width = Emu(2 * PX)
        text(s, M, y + 24, 520, 100, h, 36, ORANGE, bold=True)
        text(s, M + 560, y + 24, W - 560, 110, body, 32, BODY)

    # 10. Vision
    s = new_slide(prs, SAND, "Tracks 1 and 3 live here. The dataset already holds the material for the handoff (April close, Ellen takes over) and for onboarding (Owen joins in March).")
    title(s, "After the weekend")
    cols = [("FOR THE CONTROLLER", ["Flags controls skipped this month", "Checks history before chasing a colleague", "Rebuilds the trail behind any number"]),
            ("FOR THE TEAM", ["Sourced handoff when someone is on leave", "Onboarding from the approved procedure", "Procedure changes with re-approval"]),
            ("FOR THE CFO", ["Forecasts that already include what was said in meetings", '"Repeats next month?" on every line'])]
    cw = (W - 2 * 32) // 3
    for i, (h, lines) in enumerate(cols):
        x = M + i * (cw + 32)
        card(s, x, 290, cw, 560, LIGHT)
        text(s, x + 40, 330, cw - 80, 35, h, 24, BLUE, bold=True)
        text(s, x + 40, 390, cw - 80, 440, lines, 32, BODY, spacing=1.3)

    # 11. Team and 48 hours
    s = new_slide(prs, LIGHT, "Fill the names once the team is formed on Friday. Shows the jury the scope fits 48 hours. Confirm at the briefing that a prepared fictional dataset is allowed.")
    title(s, "The team and the 48 hours")
    half = (W - 48) // 2
    table(s, M, 260, half, [["Role", "Owns"], ["[Name] - Capture", "screenpipe setup, recording, API and MCP"],
                            ["[Name] - Engine", "Variance code, data checks, tests"], ["[Name] - Memory", "Linking variances to evidence"],
                            ["[Name] - Interface", "Variance view, sources, commentary"], ["[Name] - Story", "Pitch, video, live demo"]],
          [0.38, 0.62], size=24, row_h=72)
    table(s, M + half + 48, 260, half, [["When", "Checkpoint"], ["Fri night", "Engine matches the answer key, rejects the bad export"],
                                        ["Sat 13:00", 'Evidence for 5 of 7 variances, freight "not found"'], ["Sat 18:00", "Blockers to screenpipe office hours"],
                                        ["Sun 12:00", "Feature freeze, rehearse, record"], ["Sun 15:30", "Video, GitHub link, description submitted"]],
          [0.28, 0.72], size=24, row_h=72)
    text(s, M, 760, W, 100, "Ready before Friday: a fictional company with three months of books, mail, chat and meetings, and an answer key.", 32, BODY)

    # 12. Close
    s = new_slide(prs, DARK, "End on the outcome, then the nod to screenpipe. Pause. Then: thank you, questions.")
    text(s, M, 230, W, 300, "The controller stops searching and starts analysing.", 104, LIGHT, bold=True, spacing=1.05)
    text(s, M, 560, 1500, 130, "Our fictional company is called Pipeline. We named it after the tool that gave it a memory.", 44, ORANGE_LIGHT)
    footer(s, "Close Memory - 42 Paris x screenpipe hackathon - 16 to 18 October 2026", ON_DARK_MUTED)

    prs.core_properties.title = "Close Memory - hackathon pitch"
    prs.core_properties.author = "Close Memory team"
    prs.core_properties.last_modified_by = "Close Memory team"
    prs.core_properties.created = FIXED_TIME
    prs.core_properties.modified = FIXED_TIME
    prs.core_properties.revision = 1
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    normalize_zip(OUT)
    print("pitch deck written to", OUT)


if __name__ == "__main__":
    build()
